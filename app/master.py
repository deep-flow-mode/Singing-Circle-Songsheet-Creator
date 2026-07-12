"""Liest den Master-Katalog aus der Songtabelle.xlsx (Blatt '26')."""
import openpyxl

from .config import import_dir, MASTER_XLSX_NAME, MASTER_SHEET


def _clean(v) -> str:
    return " ".join(str(v).split()) if v is not None else ""


def load_master() -> list[dict]:
    """-> Liste von {liednummer, liedtitel, tonart}."""
    path = import_dir() / MASTER_XLSX_NAME
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[MASTER_SHEET] if MASTER_SHEET in wb.sheetnames else wb[wb.sheetnames[0]]

    songs = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        nr = row[0] if len(row) > 0 else None
        titel = row[3] if len(row) > 3 else None
        tonart = row[4] if len(row) > 4 else None
        if titel is None or not str(titel).strip():
            continue
        nr_clean = None
        if nr is not None and str(nr).strip():
            nr_clean = str(nr).strip()
            if nr_clean.endswith(".0"):
                nr_clean = nr_clean[:-2]
        songs.append({
            "liednummer": nr_clean,
            "liedtitel": _clean(titel),
            "tonart": _clean(tonart) or None,
        })
    wb.close()
    return songs
