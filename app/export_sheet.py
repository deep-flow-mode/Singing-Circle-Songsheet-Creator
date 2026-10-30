"""Erzeugt ein druckbares Singkreis-Sheet (A4, 2 Spalten) als .docx.

- Oben: Titeltext + Datum (volle Breite)
- Darunter: die Lieder in 2 Spalten
- Das Tool wählt selbst die größte Schrift (20-17 pt), die auf TARGET_PAGES Seiten passt
- Reicht das selbst mit MIN_FONT nicht: Lieder dürfen zwischen Strophen umbrechen, wenn das
  eine Seite spart (so wenige wie nötig); sonst kommen weitere Seiten hinzu
- Mehr als MAX_PAGES Seiten: NeedsMultipage wird ausgelöst (Warnhinweis im Tool)
- QR-Codes nur auf Seite 1; die Folgeseiten nutzen den Platz für Lieder
"""
from __future__ import annotations

import re
import unicodedata
from datetime import datetime

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

from .config import SHEETS_DIR
from .i18n import _, current_lang
from .qr import ensure_qr_images

PAGE_W_MM, PAGE_H_MM = 210, 297      # A4
MARGIN_MM = 12
MARGIN_BOTTOM_MM = 48                 # Seite 1: Platz unten für QR-Codes + Seitenzahl
MARGIN_BOTTOM_NEXT_MM = 16            # Folgeseiten: nur Seitenzahl
COL_SPACE_MM = 7
FONT_LADDER = [20, 19, 18, 17]        # Schrift so groß wie möglich
MIN_FONT = FONT_LADDER[-1]
TARGET_PAGES = 2                      # mehr Seiten nur, wenn es selbst bei MIN_FONT nicht passt
MAX_PAGES = 3                         # darüber nur nach Warnhinweis
MM_TO_PT = 72 / 25.4
LINE_H = 1.22                         # Zeilenhöhe von Calibri bei einfachem Abstand
SONG_GAP = 0.6                        # Abstand vor einem Lied, relativ zur Schrift
FILL = 0.97                           # Reserve, damit Word nicht knapp überläuft
FOOTER_PT = 8
GREY = RGBColor(0x88, 0x88, 0x88)
QR_WIDTH_MM = 27


class NeedsMultipage(Exception):
    """Wird ausgelöst, wenn der Inhalt selbst bei Mindestschrift mehr als MAX_PAGES Seiten braucht."""
    def __init__(self, pages: int):
        super().__init__(f"benötigt {pages} Seiten")
        self.pages = pages


def _set_columns(section, num: int, space_mm: float = COL_SPACE_MM) -> None:
    sectPr = section._sectPr
    cols = sectPr.find(qn("w:cols"))
    if cols is None:
        cols = OxmlElement("w:cols")
        sectPr.append(cols)
    cols.set(qn("w:num"), str(num))
    cols.set(qn("w:space"), str(int(space_mm * 56.7)))  # mm -> twips


def _setup_section(section, bottom_mm: float) -> None:
    section.page_width = Mm(PAGE_W_MM)
    section.page_height = Mm(PAGE_H_MM)
    section.top_margin = Mm(MARGIN_MM)
    section.left_margin = Mm(MARGIN_MM)
    section.right_margin = Mm(MARGIN_MM)
    section.bottom_margin = Mm(bottom_mm)
    section.footer_distance = Mm(6)


# Zeichenbreiten von Calibri (Anteil der Schriftgröße) für die Umbruch-Berechnung
_CHAR_W = {
    "a": .479, "b": .525, "c": .423, "d": .525, "e": .498, "f": .305, "g": .471,
    "h": .525, "i": .229, "j": .239, "k": .455, "l": .229, "m": .799, "n": .525,
    "o": .527, "p": .525, "q": .525, "r": .349, "s": .391, "t": .335, "u": .525,
    "v": .452, "w": .715, "x": .433, "y": .453, "z": .395, " ": .226, ".": .252,
    ",": .25, "!": .326, "?": .463, "-": .306, "'": .221, "(": .303, ")": .303,
}


def _text_width(text: str, font_pt: float, bold: bool = False) -> float:
    width = 0.0
    for ch in text:
        base = unicodedata.normalize("NFD", ch)[0]
        if base.isupper():
            width += .6
        elif base.isdigit():
            width += .507
        else:
            width += _CHAR_W.get(base, .5)
    return width * font_pt * (1.03 if bold else 1.0)


def _line_count(line: str, font_pt: float, width_pt: float, bold: bool = False) -> int:
    """Wie viele Zeilen belegt `line`, wenn sie an Wortgrenzen umbricht?"""
    lines, used = 1, 0.0
    space = _text_width(" ", font_pt, bold)
    for word in line.split(" "):
        w = _text_width(word, font_pt, bold)
        if used and used + space + w > width_pt:
            lines += 1
            used = 0.0
        if used:
            used += space
        while w > width_pt:                    # sehr langes Wort bricht mitten drin
            lines += 1
            w -= width_pt
        used += w
    return lines


def _title_pt(body_pt: float) -> float:
    return max(20, round(body_pt * 1.4))


def _date_pt(body_pt: float) -> float:
    return max(12, round(body_pt * 0.8))


def _title_block_height(titel: str, datum: str, font_pt: float) -> float:
    width = (PAGE_W_MM - 2 * MARGIN_MM) * MM_TO_PT
    title_pt = _title_pt(font_pt)
    height = _line_count(titel or _("Singkreis"), title_pt, width, bold=True) * title_pt * LINE_H
    if datum:
        height += font_pt * LINE_H  # Absatzmarke hat Grundschrift -> bestimmt die Zeilenhöhe
    return height + 2               # Abschnittswechsel


def _clean_text(song: dict) -> str:
    # \\r entfernen: python-docx macht daraus sonst zusätzliche Umbrüche (Leerzeilen)
    return (song.get("liedtext") or "").replace("\r\n", "\n").replace("\r", "\n").strip()


def _stanzas(text: str) -> list:
    return re.split(r"\n\s*\n", text) if text else []


def _column_height(bottom_mm: float = MARGIN_BOTTOM_NEXT_MM) -> float:
    return (PAGE_H_MM - MARGIN_MM - bottom_mm) * MM_TO_PT * FILL


def _text_height(text: str, font_pt: float) -> float:
    width = ((PAGE_W_MM - 2 * MARGIN_MM - COL_SPACE_MM) / 2) * MM_TO_PT
    return sum(_line_count(ln, font_pt, width) for ln in text.split("\n")) * font_pt * LINE_H


def _song_blocks(song: dict, font_pt: float, first: bool, split: bool = False) -> list:
    """Höhen der Blöcke, die zusammenbleiben müssen: normal das ganze Lied. Geteilt in
    Strophen wird nur auf Wunsch (`split`) oder wenn es länger als eine ganze Spalte ist."""
    width = ((PAGE_W_MM - 2 * MARGIN_MM - COL_SPACE_MM) / 2) * MM_TO_PT
    title_pt = font_pt + 1
    head = 0.0 if first else font_pt * SONG_GAP
    head += _line_count(song["liedtitel"], title_pt, width, bold=True) * title_pt * LINE_H + 1
    heights = [_text_height(st, font_pt) for st in _stanzas(_clean_text(song))] or [0.0]
    spacer = font_pt * LINE_H
    total = head + sum(heights) + spacer * (len(heights) - 1)
    if total <= _column_height() and not split:
        return [total]
    return [head + heights[0]] + [spacer + h for h in heights[1:]]


def _layout(songs, titel: str, datum: str, font_pt: float, splittable=frozenset()):
    """Setzt die Lieder probeweise Spalte für Spalte wie Word: ein Block, der nicht mehr
    in die Spalte passt, rutscht komplett in die nächste. Lieder in `splittable` dürfen
    stattdessen zwischen Strophen umbrechen. -> (seiten, umgebrochene Lieder)"""
    column = _column_height()
    first = _column_height(MARGIN_BOTTOM_MM) - _title_block_height(titel, datum, font_pt)

    def capacity(col: int) -> float:
        return first if col < 2 else column   # Seite 1: Titel oben, QR-Codes unten

    col, used, broken = 0, 0.0, set()
    for i, song in enumerate(songs):
        blocks = _song_blocks(song, font_pt, first=i == 0)
        if i in splittable and len(blocks) == 1 and used + blocks[0] > capacity(col):
            parts = _song_blocks(song, font_pt, first=i == 0, split=True)
            if len(parts) > 1 and used + parts[0] <= capacity(col):
                blocks = parts
                broken.add(i)
        for height in blocks:
            if used and used + height > capacity(col):
                col, used = col + 1, 0.0
            while height > capacity(col) and capacity(col) < column:
                col += 1                       # passt nur in eine Spalte ohne Titel
            while height > column:             # einzelne Strophe länger als eine Spalte
                height -= column
                col += 1
            used += height
    return col // 2 + 1, broken


def choose_font(songs, titel: str = "", datum: str = ""):
    """-> (schriftgröße, seiten, umgebrochene Lieder). Größte Schrift, die auf TARGET_PAGES
    Seiten passt. Reicht selbst die Mindestschrift nicht, dürfen Lieder zwischen Strophen
    umbrechen, aber nur wenn das eine Seite spart; sonst gilt die nötige Seitenzahl."""
    def pages_at(font_pt, allowed=frozenset()):
        return _layout(songs, titel, datum, font_pt, allowed)

    limit, allowed = pages_at(MIN_FONT)[0], frozenset()
    if limit > TARGET_PAGES:
        everyone = frozenset(range(len(songs)))
        with_breaks = pages_at(MIN_FONT, everyone)[0]
        if with_breaks < limit:
            limit, allowed = with_breaks, everyone
    limit = max(limit, TARGET_PAGES)

    for font_pt in FONT_LADDER:
        pages, broken = pages_at(font_pt, allowed)
        if pages <= limit:
            # So wenige Umbrüche wie möglich: jeden einzeln wieder wegnehmen, wenn es noch passt
            for i in sorted(broken, reverse=True):
                if pages_at(font_pt, frozenset(broken - {i}))[0] <= limit:
                    broken = broken - {i}
            return (font_pt,) + pages_at(font_pt, frozenset(broken))
    return (MIN_FONT,) + pages_at(MIN_FONT, allowed)


def _fmt_date(d: str) -> str:
    try:
        fmt = "%d %b %Y" if current_lang() == "en" else "%d.%m.%Y"
        return datetime.strptime(d, "%Y-%m-%d").strftime(fmt)
    except Exception:
        return d or ""


def _safe(s: str) -> str:
    s = re.sub(r"[^\w\- ]", "", s or "", flags=re.UNICODE).strip().replace(" ", "_")
    return s or "Sheet"


def _text_run(paragraph, text, size):
    r = paragraph.add_run(text)
    r.font.size = Pt(size)
    r.font.color.rgb = GREY
    return r


def _field_run(paragraph, code, size):
    """Fügt ein Word-Feld (z.B. PAGE / NUMPAGES) als Run ein."""
    run = paragraph.add_run()
    run.font.size = Pt(size)
    run.font.color.rgb = GREY
    begin = OxmlElement("w:fldChar"); begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText"); instr.set(qn("xml:space"), "preserve"); instr.text = code
    end = OxmlElement("w:fldChar"); end.set(qn("w:fldCharType"), "end")
    run._r.append(begin)
    run._r.append(instr)
    run._r.append(end)
    return run


def _build_footer(footer, qrs=()):
    """Fußbereich: optional QR-Codes (mit Beschriftung) + zentrierte Seitenzahl."""
    footer.is_linked_to_previous = False
    # Standardinhalt der Fußzeile entfernen
    for p in list(footer.paragraphs):
        p._element.getparent().remove(p._element)

    if qrs:
        # Spalten über die volle Nutzbreite verteilen -> Codes weit auseinander
        col_mm = (PAGE_W_MM - 2 * MARGIN_MM) / len(qrs)
        table = footer.add_table(rows=2, cols=len(qrs), width=Mm(col_mm * len(qrs)))
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        for i, (label, path) in enumerate(qrs):
            table.columns[i].width = Mm(col_mm)
            cell_img = table.cell(0, i)
            cell_img.width = Mm(col_mm)
            pi = cell_img.paragraphs[0]
            pi.style = "Footer"
            pi.alignment = WD_ALIGN_PARAGRAPH.CENTER
            pi.add_run().add_picture(str(path), width=Mm(QR_WIDTH_MM))
            cell_lab = table.cell(1, i)
            cell_lab.width = Mm(col_mm)
            pl = cell_lab.paragraphs[0]
            pl.style = "Footer"
            pl.alignment = WD_ALIGN_PARAGRAPH.CENTER
            rl = pl.add_run(label)
            rl.font.size = Pt(FOOTER_PT)
            rl.font.color.rgb = GREY

    pn = footer.add_paragraph(style="Footer")
    pn.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _text_run(pn, _("Seite "), FOOTER_PT)
    _field_run(pn, "PAGE", FOOTER_PT)
    _text_run(pn, _(" von "), FOOTER_PT)
    _field_run(pn, "NUMPAGES", FOOTER_PT)


def create_sheet(songs, titel: str, datum: str, allow_multipage: bool = False):
    """-> (pfad, schriftgröße, seiten). Löst NeedsMultipage aus, wenn nötig."""
    body_font, pages, broken = choose_font(songs, titel, datum)
    if pages > MAX_PAGES and not allow_multipage:
        raise NeedsMultipage(pages)

    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(body_font)
    normal.paragraph_format.space_after = Pt(0)
    normal.paragraph_format.line_spacing = 1.0
    # Fußzeile unabhängig von der Grundschrift, sonst wüchse sie mit und nähme den Liedern Platz
    doc.styles["Footer"].font.size = Pt(FOOTER_PT)

    # Abschnitt 1: volle Breite -> Titel + Datum
    sec0 = doc.sections[0]
    _setup_section(sec0, MARGIN_BOTTOM_MM)
    _set_columns(sec0, 1)

    p_titel = doc.add_paragraph()
    p_titel.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p_titel.add_run(titel or _("Singkreis"))
    r.bold = True
    r.font.size = Pt(_title_pt(body_font))
    r.font.color.rgb = RGBColor(0x3C, 0x4A, 0x2E)

    if datum:
        p_datum = doc.add_paragraph()
        p_datum.alignment = WD_ALIGN_PARAGRAPH.CENTER
        rd = p_datum.add_run(_fmt_date(datum))
        rd.font.size = Pt(_date_pt(body_font))
        rd.font.color.rgb = RGBColor(0x66, 0x66, 0x66)

    # Abschnitt 2: fortlaufend, 2 Spalten -> Lieder
    sec1 = doc.add_section(WD_SECTION.CONTINUOUS)
    # Der Abschnittswechsel steckt in einem eigenen leeren Absatz -> keine Leerzeile erzeugen
    doc.paragraphs[-1].paragraph_format.line_spacing = Pt(2)
    _setup_section(sec1, MARGIN_BOTTOM_NEXT_MM)
    _set_columns(sec1, 2)

    for i, s in enumerate(songs):
        pt = doc.add_paragraph()
        pt.paragraph_format.space_before = Pt(0 if i == 0 else body_font * SONG_GAP)
        pt.paragraph_format.space_after = Pt(1)
        pt.paragraph_format.keep_with_next = True
        pt.paragraph_format.keep_together = True
        rt = pt.add_run(s["liedtitel"])
        rt.bold = True
        rt.font.size = Pt(body_font + 1)

        text = _clean_text(s)
        # Ein Lied bleibt ganz, außer es spart eine Seite oder ist länger als eine Spalte
        whole = i not in broken and len(_song_blocks(s, body_font, first=i == 0)) == 1
        parts = [text] if whole else _stanzas(text)
        for k, part in enumerate(filter(None, parts)):
            pb = doc.add_paragraph()
            pb.paragraph_format.keep_together = _text_height(part, body_font) <= _column_height()
            if k < len(parts) - 1:
                pb.paragraph_format.space_after = Pt(body_font * LINE_H)
            for j, line in enumerate(part.split("\n")):
                if j:
                    pb.add_run().add_break()
                pb.add_run(line)

    # Seite 1 gehört zum ersten Abschnitt: eigene Fußzeile mit QR-Codes.
    # Folgeseiten gehören zum zweiten Abschnitt: nur die Seitenzahl.
    first_sec, next_sec = doc.sections[0], doc.sections[1]  # neu holen: add_section verschiebt die Bezüge
    first_sec.different_first_page_header_footer = True
    _build_footer(first_sec.first_page_footer, ensure_qr_images())
    _build_footer(first_sec.footer)
    next_sec.different_first_page_header_footer = False
    _build_footer(next_sec.footer)

    SHEETS_DIR.mkdir(parents=True, exist_ok=True)
    base = f"{_safe(datum)}_{_safe(titel)}".strip("_")
    path = SHEETS_DIR / f"{base}.docx"
    n = 2
    while path.exists():
        path = SHEETS_DIR / f"{base}_{n}.docx"
        n += 1
    doc.save(str(path))
    return path, body_font, pages
